"""Quince Matrix daily job.
Usage: python3 daily.py <export.xlsx> <base LIVE_ODR workbook.xlsx> [YYYY-MM-DD]
Run from the engine folder (/home/claude/build). Writes:
  state.json (updated in place), quince-matrix.html (private dashboard),
  site/index.html (public, no AWBs/order numbers), out/Quince_Matrix_Shop.xlsx,
  out/LIVE_ODR_QTY - <tab>.xlsx ; prints a JSON summary on the last line.
"""
import shutil, sys, os, json, glob, hashlib, datetime, subprocess, copy, collections
import openpyxl

B = os.path.dirname(os.path.abspath(__file__)) + '/'
os.makedirs(B + 'out', exist_ok=True); os.makedirs(B + 'site', exist_ok=True)
exp_path, base_odr = sys.argv[1], sys.argv[2]
ist = datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=5, minutes=30)))
day = sys.argv[3] if len(sys.argv) > 3 else ist.date().isoformat()
D = datetime.date.fromisoformat(day)

def ordinal(n):
    return f"{n}{'th' if 11 <= n % 100 <= 13 else {1: 'st', 2: 'nd', 3: 'rd'}.get(n % 10, 'th')}"
TAB = f"{ordinal(D.day)} {D.strftime('%b').lower()}"          # e.g. "20th oct"
TITLE = f"{ordinal(D.day)} {D.strftime('%b')}"                 # e.g. "20th Oct"

# ---------- read export (by column name, so new portal columns do not break it) ----------
ew = openpyxl.load_workbook(exp_path).worksheets[0]
hdr = [str(c.value).strip() if c.value is not None else '' for c in ew[1]]
rows = [dict(zip(hdr, r)) for r in ew.iter_rows(min_row=2, values_only=True) if any(v not in (None, '') for v in r)]
need = ['SKU', 'Qty', 'Variant', 'Shipping Destination', 'lastMileTrackingNumber', 'orderNumber']
missing = [k for k in need if k not in hdr]
if missing:
    print(json.dumps({'ok': False, 'error': 'export columns missing: ' + ', '.join(missing)})); sys.exit(2)

# ---------- update shop state ----------
st = json.load(open(B + 'state.json'))
known = {s[4] for s in st['sales'] if s[4]}
cat = {c['sku'] for c in st['catalog']}
new, skipped, unknown = [], 0, []
for x in rows:
    awb = str(x['lastMileTrackingNumber'] or '').strip()
    if awb and awb in known: skipped += 1; continue
    sku = str(x['SKU'] or '').strip().upper()
    if sku not in cat: unknown.append(sku); continue
    dest = str(x['Shipping Destination'] or '').split(',')[-1].strip().upper()
    line = [day, str(x['orderNumber']).strip(), sku, int(float(x['Qty'] or 1)), awb, 'CA' if dest == 'CA' else 'US' if dest == 'US' else '']
    st['sales'].append(line); known.add(awb); new.append((x, line))
pcs = sum(l[3] for _, l in new)
if new:
    st['updated'] = datetime.datetime.now(datetime.timezone.utc).isoformat(timespec='seconds')
    st['log'].append({'at': day, 'kind': 'orders', 'f': os.path.basename(exp_path), 'pcs': pcs,
                      'n': f"{len(new)} new lines for {D.strftime('%d %b %Y')}, {skipped} already counted" + (f", {len(unknown)} unknown SKU" if unknown else '')})
    json.dump(st, open(B + 'state.json', 'w'), separators=(',', ':'))

# ---------- balances for the summary ----------
insp = collections.Counter(); sold = collections.Counter()
B2B = set((st.get('b2b') or {}).get('pos', []))
for l in st['insp']:
    if l['po'] not in B2B: insp[l['s']] += l['q']
for s in st['sales']: sold[s[2]] += s[3]
red, yel = st['settings']['red'], st['settings']['yellow']
def status(b): return 'empty' if b <= 0 else 'red' if b <= red else 'yellow' if b <= yel else 'green'
today_skus = collections.Counter(l[2] for _, l in new)
moved = {k: {'sold_today': v, 'stock': insp[k] - sold[k], 'status': status(insp[k] - sold[k])} for k, v in today_skus.items()}
oversold = [k for k in insp if insp[k] - sold[k] < 0] + [k for k in sold if k not in insp]

# ---------- private dashboard + public site ----------
subprocess.run([sys.executable, '-I', B + 'assemble.py', B + 'state.json', B + 'quince-matrix.html'], check=True, stdout=subprocess.DEVNULL)
h = lambda o: hashlib.sha256(('qm-' + str(o).strip()).encode()).hexdigest()[:10]
pub = dict(st); pub['sales'] = [[d, h(o), s, q, '', c] for d, o, s, q, a, c in st['sales']]
json.dump(pub, open(B + 'out/public_state.json', 'w'), separators=(',', ':'))
subprocess.run([sys.executable, '-I', B + 'assemble.py', B + 'out/public_state.json', B + 'site/index.html'], check=True, stdout=subprocess.DEVNULL)
s = open(B + 'site/index.html').read()
s = s.replace('Customer names and addresses are not stored on this page.',
              "Order numbers, AWB numbers and customer details are not published on this page. Updated after each day's order confirmation.")
# production (separate from stock): Firebase scripts on the public page + PO catalog for production.html
FB = 'https://www.gstatic.com/firebasejs/10.12.2/firebase-'
tags = ''.join('<script src="%s%s-compat.js"></script>' % (FB, m) for m in ('app', 'auth', 'firestore')) + '<script src="qm-fire.js"></script><script src="prod-tab.js"></script>'
s = s.rstrip() + '\n' + tags + '\n'  # page has no </body>; the inline JS contains '</body>' text, so append at the end
if not s.lstrip().lower().startswith('<!doctype'): s = '<!doctype html>\n<meta charset="utf-8">\n' + s
open(B + 'site/index.html', 'w').write(s)
for f in ('production.html', 'qm-fire.js', 'prod-tab.js', 'fabric.html', 'qm-fabric.js'):
    shutil.copy(B + 'web/' + f, B + 'site/' + f)
inspq = collections.Counter()
for l in st['insp']: inspq[(l['po'], l['s'])] += l['q']
pos_c = {}
for c in st['catalog']:
    for p in c.get('pos', []):
        P = pos_c.setdefault(p['po'], {'po': p['po'], 'lines': [], 'qty': 0})
        P['lines'].append({'sku': c['sku'], 'product': c['product'], 'colour': c['colour'], 'size': c['size'], 'cls': c['cls'], 'qty': p['qty'], 'insp': inspq[(p['po'], c['sku'])]})
        P['qty'] += p['qty']
ship = st.get('b2b', {}).get('ship', {})
for P in pos_c.values():
    P['b2b'] = P['po'] in B2B
    P['ship'] = ship.get(str(P['po']), '')
    P['products'] = sorted({l['product'] for l in P['lines']})
json.dump({'updated': st['updated'], 'fabric_history': st.get('fabric_history'), 'pos': sorted(pos_c.values(), key=lambda P: -P['po'])}, open(B + 'site/catalog.json', 'w'), separators=(',', ':'))
os.remove(B + 'out/public_state.json')

# ---------- Excel shop ----------
xl = B + 'out/Quince_Matrix_Shop.xlsx'
subprocess.run([sys.executable, '-I', B + 'excel2.py', xl], check=True, stdout=subprocess.DEVNULL, cwd=B)
rc = sorted(glob.glob('/root/.claude/skills/**/xlsx/scripts/recalc.py', recursive=True))
recalc = 'skipped'
if rc:
    r = subprocess.run([sys.executable, rc[0], xl, '400'], capture_output=True, text=True, cwd=os.path.dirname(os.path.dirname(rc[0])), timeout=900)
    recalc = 'ok' if '"status": "success"' in r.stdout else ('errors: ' + r.stdout[-300:])

# ---------- LIVE ODR QTY: add today's tab ----------
odr_out = None
if new:
    wb = openpyxl.load_workbook(base_odr)
    if TAB in wb.sheetnames: del wb[TAB]
    tpl = None
    for ws in reversed(wb.worksheets):
        vals = [c.value for c in ws[2]] if ws.max_row >= 2 else []
        if 'S NO.' in vals and 'CUST NAME' in vals: tpl = ws; break
    ws = wb.copy_worksheet(tpl); ws.title = TAB
    for m in list(ws.merged_cells.ranges): ws.unmerge_cells(str(m))
    # template styles: section row = first text row after header, data row = first numbered row
    sec_r = next(r for r in range(3, tpl.max_row + 1) if isinstance(tpl.cell(r, 1).value, str))
    dat_r = next(r for r in range(3, tpl.max_row + 1) if isinstance(tpl.cell(r, 1).value, (int, float)))
    sec_style = [copy.copy(tpl.cell(sec_r, c)._style) for c in range(1, 10)]
    data_style = [copy.copy(tpl.cell(dat_r, c)._style) for c in range(1, 10)]
    blank = copy.copy(tpl.cell(sec_r, 12)._style)
    for r in range(3, ws.max_row + 1):
        for c in range(1, 14): ws.cell(r, c).value = None; ws.cell(r, c)._style = copy.copy(blank)
    ws['A1'] = f'LIST OF CONFIRMED ORDERS ({TITLE})'
    def size(v):
        t = str(v).split('/', 1)[-1].strip().upper()
        return '50X70' if '50' in t else 'TWIN' if 'TWIN' in t else 'FULL/QUEEN' if 'QUEEN' in t else 'KING/CAL KING' if 'KING' in t else t
    def conv(x, l):
        awb = l[4]
        return [str(x['Variant']).split('/')[0].strip().upper(), int(l[1]) if l[1].isdigit() else l[1], size(x['Variant']), l[3],
                x.get('Customer') or '', int(awb) if awb.isdigit() else awb]
    ca = sorted([conv(x, l) for x, l in new if l[5] == 'CA'], key=lambda v: (str(v[1]), str(v[5])))
    us = sorted([conv(x, l) for x, l in new if l[5] != 'CA'], key=lambda v: (str(v[1]), str(v[5])))
    r = 3; n = 1
    for title, group in [("TODAY'S CANADA ORDERS", ca), ("TODAY'S USA ORDERS", us)]:
        ws.cell(r, 1, title)
        for c in range(1, 10): ws.cell(r, c)._style = copy.copy(sec_style[c - 1])
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9); r += 1
        for v in group:
            for c, val in enumerate([n] + v + [None, None], 1):
                ws.cell(r, c, val)._style = copy.copy(data_style[c - 1])
            r += 1; n += 1
    ws.merge_cells('A1:I1')
    bold = copy.copy(tpl.cell(2, 1).font)
    ws.cell(r, 4, 'TOTAL').font = bold; ws.cell(r, 5, f'=SUM(E4:E{r-1})').font = bold
    for w in wb.worksheets: w.sheet_view.tabSelected = (w.title == TAB)
    wb.active = wb.sheetnames.index(TAB)
    odr_out = B + f'out/LIVE_ODR_QTY - {TAB}.xlsx'
    wb.save(odr_out)

ins_total = sum(insp.values()); sold_total = sum(sold.values())
print(json.dumps({'ok': True, 'day': day, 'tab': TAB, 'export_lines': len(rows), 'new_lines': len(new), 'new_pcs': pcs,
                  'new_orders': len({l[1] for _, l in new}), 'already_counted': skipped, 'unknown_skus': unknown,
                  'canada': sum(1 for _, l in new if l[5] == 'CA'), 'moved': moved, 'oversold': oversold,
                  'inspected': ins_total, 'sold': sold_total, 'in_stock': sum(max(insp[k] - sold[k], 0) for k in insp),
                  'b2b_pending': sum(l['q'] for l in st['insp'] if l['po'] in B2B) - sum(d['q'] for d in (st.get('b2b') or {}).get('dispatch', [])),
                  'excel_recalc': recalc, 'odr_file': odr_out}))
